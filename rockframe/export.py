# ROCKFRAME – Exportación de resultados a Excel (sin interfaz gráfica)

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from rockframe.engine import N_to_kN

COLUMNAS_FUERZAS = ["elem", "Ni[kN]", "Vi[kN]", "Mi[kN·m]", "Nj[kN]", "Vj[kN]", "Mj[kN·m]"]
COLUMNAS_REACCIONES = ["nodo", "Rx[kN]", "Ry[kN]", "M[kN·m]"]


def _fuerzas(filas):
    return pd.DataFrame([(eid, *(N_to_kN(v) for v in f)) for (eid, *f) in filas], columns=COLUMNAS_FUERZAS)


def _reacciones(reacciones):
    return pd.DataFrame(
        [(nid, N_to_kN(Fx), N_to_kN(Fy), N_to_kN(Mz)) for nid, (Fx, Fy, Mz) in sorted(reacciones.items())],
        columns=COLUMNAS_REACCIONES,
    )


def exportar_excel(path, resumen_modelo, espectro, estatico, sismico):
    """resumen_modelo y espectro son diccionarios {parámetro: valor}."""
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"Parámetro": list(resumen_modelo), "Valor": list(resumen_modelo.values())}).to_excel(
            xw, sheet_name="0_RESUMEN_MODELO", index=False)
        pd.DataFrame({"Parámetro": list(espectro), "Valor": list(espectro.values())}).to_excel(
            xw, sheet_name="espectro", index=False)

        s = sismico
        pd.DataFrame([{
            "Criterio_periodo": s["periodo_diseno"],
            "Ta [s]": s["T_aprox_inicial"],
            "Cu": s["Cu"],
            "Cu·Ta [s]": s["T_limite"],
            "T_Rayleigh [s]": s["T_detallado"],
            "T_diseno [s]": s["T_diseno"],
            "Sa [g]": s["Sa"],
            "k": s["k_exp"],
            "W_total [kN]": s["W_total"],
            "V_basal [kN]": s["V"],
            "Iteraciones": s["iteration_count"],
            "Convergió": "Sí" if s["converged"] else "No",
            "Deriva_max_calculada [%]": max(s["drift_pct"]),
            "Deriva_max_evaluada [%]": s["drift_max_evaluacion_pct"],
            "Limite_deriva [%]": s["limite_deriva_pct"],
            "Cumple_deriva": "Sí" if s["cumple_deriva"] else "No",
            "Q_max": s["Q_max"],
            "Carga_viva_en_Q": "Sí" if s["carga_viva_incluida"] else "No",
        }]).T.rename(columns={0: "Valor"}).to_excel(xw, sheet_name="sismo_resumen", index_label="Parámetro")
        pd.DataFrame([{
            "Piso": k + 1,
            "Altura [m]": s["hi"][k],
            "Wi [kN]": s["Wi"][k],
            "Fi [kN]": s["Fi"][k],
            "Desplazamiento [m]": s["disp_piso"][k],
            "Deriva [%]": s["drift_pct"][k],
            "Deriva evaluada [%]": s["drift_evaluacion_pct"][k],
            "Qi": s["Qi"][k],
        } for k in range(s["N_pisos"])]).to_excel(xw, sheet_name="pisos_sismo", index=False)
        pd.DataFrame([{
            "Iteración": it["iter"],
            "T_entrada [s]": it["T_in"],
            "Sa(T) [g]": it["Sa"],
            "k": it["k_exp"],
            "V [kN]": it["V"],
            "T_Rayleigh [s]": it["T_rayleigh"],
            "Error relativo [%]": it["error_pct"],
            "Convergió": "Sí" if it["converged"] else "No",
        } for it in s["iterations"]]).to_excel(xw, sheet_name="iteraciones", index=False)
        _fuerzas(s["elem_forces_local"]).to_excel(xw, sheet_name="esf_sismicos", index=False)
        _reacciones(s["reactions"]).to_excel(xw, sheet_name="reacciones_sismicas", index=False)

        e = estatico
        _fuerzas(e["elem_forces_local"]).to_excel(xw, sheet_name="esf_estaticos", index=False)
        _reacciones(e["reactions_by_node"]).to_excel(xw, sheet_name="reacciones_estaticas", index=False)
        pd.DataFrame([(nid, n.x, n.y) for nid, n in sorted(e["nodes"].items())],
                     columns=["nodo", "x[m]", "y[m]"]).to_excel(xw, sheet_name="nodos", index=False)
        pd.DataFrame([(x.id, x.i, x.j, x.A, x.I, x.E, x.w_kNpm) for x in e["elements"].values()],
                     columns=["elem", "i", "j", "A[m2]", "I_bruta[m4]", "E[Pa]", "w[kN/m]"]).to_excel(
            xw, sheet_name="elementos", index=False)
        filas_U = []
        for nid in sorted(e["nodes"]):
            ux, uy, th = e["dof"].node_dofs[nid]
            filas_U.append((nid, e["U"][ux], e["U"][uy], e["U"][th], s["U"][ux], s["U"][uy], s["U"][th]))
        pd.DataFrame(filas_U, columns=["nodo", "ux_est[m]", "uy_est[m]", "giro_est[rad]",
                                       "ux_sis[m]", "uy_sis[m]", "giro_sis[rad]"]).to_excel(
            xw, sheet_name="desplazamientos", index=False)
        pd.DataFrame(e["K"]).to_excel(xw, sheet_name="K_global", index=False)

    wb = load_workbook(path)
    for nombre in wb.sheetnames:
        ws = wb[nombre]
        if nombre == "K_global":
            continue
        for celda in ws[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill(start_color="1C2326", end_color="1C2326", fill_type="solid")
        for columna in ws.columns:
            letra = columna[0].column_letter
            ws.column_dimensions[letra].width = max(12, min(34, max(len(str(c.value or "")) for c in columna[:40]) + 2))
    wb.save(path)
